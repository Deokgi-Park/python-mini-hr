"""
Python Mini HR - 개인용 HR 관리 프로그램
간단한 직원 정보 관리 시스템
"""

from flask import Flask, render_template, request, redirect, url_for, flash, send_file
from datetime import datetime, timedelta
from dateutil.relativedelta import relativedelta
import json
import os
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO

app = Flask(__name__)
app.secret_key = 'your-secret-key-change-this-in-production'

# 데이터 저장 파일
DATA_FILE = 'employees.json'
SALARY_FILE = 'salaries.json'

# 직원 데이터 로드
def load_employees():
    if os.path.exists(DATA_FILE):
        with open(DATA_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return []

# 직원 데이터 저장
def save_employees(employees):
    with open(DATA_FILE, 'w', encoding='utf-8') as f:
        json.dump(employees, f, ensure_ascii=False, indent=2)

# 급여 데이터 로드
def load_salaries():
    if os.path.exists(SALARY_FILE):
        with open(SALARY_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return []

# 급여 데이터 저장
def save_salaries(salaries):
    with open(SALARY_FILE, 'w', encoding='utf-8') as f:
        json.dump(salaries, f, ensure_ascii=False, indent=2)

# 남은 연차 계산
def calculate_remaining_leave(annual_leave, used_leave):
    return (annual_leave or 0) - (used_leave or 0)

# 한국 근로기준법에 따른 연차 자동 계산
def calculate_annual_leave_kr(hire_date_str):
    """
    한국 근로기준법 기준 연차 계산
    - 1년 미만: 월 1개 (최대 11개)
    - 1년 이상: 15개
    - 3년 이상: 2년마다 1개씩 추가 (최대 25개)
    """
    if not hire_date_str:
        return 15  # 기본값
    
    try:
        hire_date = datetime.strptime(hire_date_str, '%Y-%m-%d')
        today = datetime.now()
        
        # 근속 기간 계산
        years = relativedelta(today, hire_date).years
        months = relativedelta(today, hire_date).months
        total_months = years * 12 + months
        
        # 1년 미만 (11개월 이하)
        if years < 1:
            # 월 1개씩 (입사 후 1개월마다)
            return min(total_months, 11)
        
        # 1년 이상 기본 15개
        annual_leave = 15
        
        # 3년 이상부터 2년마다 1개 추가
        if years >= 3:
            additional_years = years - 1  # 2년차부터 계산
            additional_leave = additional_years // 2
            annual_leave += additional_leave
        
        # 최대 25개 제한
        return min(annual_leave, 25)
    
    except Exception as e:
        print(f"연차 계산 오류: {e}")
        return 15  # 오류 시 기본값

# 메인 페이지 - 직원 목록
@app.route('/')
def index():
    employees = load_employees()
    # 남은 연차 계산하여 추가
    for emp in employees:
        emp['remaining_leave'] = calculate_remaining_leave(
            emp.get('annual_leave', 0), 
            emp.get('used_leave', 0)
        )
    return render_template('index.html', employees=employees)

# 직원 추가 페이지
@app.route('/add', methods=['GET', 'POST'])
def add_employee():
    if request.method == 'POST':
        employees = load_employees()
        
        # 새 직원 ID 생성
        new_id = max([emp.get('id', 0) for emp in employees], default=0) + 1
        
        hire_date = request.form.get('hire_date')
        
        # 입사일로부터 근속연수 자동 계산
        if hire_date:
            hire_dt = datetime.strptime(hire_date, '%Y-%m-%d')
            years_of_service = relativedelta(datetime.now(), hire_dt).years
        else:
            years_of_service = int(request.form.get('years_of_service') or 0)
        
        # 입사일 기준으로 연차 자동 계산 (한국 근로기준법)
        annual_leave = calculate_annual_leave_kr(hire_date) if hire_date else int(request.form.get('annual_leave') or 15)
        
        # 새 직원 정보 생성
        new_employee = {
            'id': new_id,
            'name': request.form.get('name'),
            'ssn': request.form.get('ssn'),
            'department': request.form.get('department'),
            'position': request.form.get('position'),
            'hire_date': hire_date,
            'resignation_date': request.form.get('resignation_date') or None,
            'address': request.form.get('address'),
            'personal_email': request.form.get('personal_email'),
            'work_email': request.form.get('work_email'),
            'personal_phone': request.form.get('personal_phone'),
            'work_phone': request.form.get('work_phone'),
            'has_insurance': request.form.get('has_insurance') == 'on',
            'salary_bank': request.form.get('salary_bank'),
            'salary_account': request.form.get('salary_account'),
            'years_of_service': years_of_service,
            'annual_leave': annual_leave,
            'used_leave': int(request.form.get('used_leave') or 0),
            'created_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        }
        
        employees.append(new_employee)
        save_employees(employees)
        
        flash('직원이 성공적으로 추가되었습니다.', 'success')
        return redirect(url_for('index'))
    
    return render_template('add_employee.html')

# 직원 상세 정보
@app.route('/employee/<int:employee_id>')
def employee_detail(employee_id):
    employees = load_employees()
    employee = next((emp for emp in employees if emp['id'] == employee_id), None)
    
    if employee is None:
        flash('직원을 찾을 수 없습니다.', 'error')
        return redirect(url_for('index'))
    
    # 남은 연차 계산
    employee['remaining_leave'] = calculate_remaining_leave(
        employee.get('annual_leave', 0), 
        employee.get('used_leave', 0)
    )
    
    # 해당 직원의 급여 정보 로드
    salaries = load_salaries()
    employee_salaries = [s for s in salaries if s['employee_id'] == employee_id]
    
    # 연도 목록 생성 (급여 데이터가 있는 연도들)
    years = sorted(list(set([s['year'] for s in employee_salaries])), reverse=True)
    
    # 선택된 연도 (기본값: 가장 최근 연도 또는 현재 연도)
    selected_year = request.args.get('year', type=int)
    if not selected_year:
        selected_year = years[0] if years else datetime.now().year
    
    # 선택된 연도의 급여 정보
    salary_info = next((s for s in employee_salaries if s['year'] == selected_year), None)
    
    return render_template('employee_detail.html', 
                         employee=employee, 
                         salary_info=salary_info,
                         years=years,
                         selected_year=selected_year)

# 직원 수정
@app.route('/edit/<int:employee_id>', methods=['GET', 'POST'])
def edit_employee(employee_id):
    employees = load_employees()
    employee = next((emp for emp in employees if emp['id'] == employee_id), None)
    
    if employee is None:
        flash('직원을 찾을 수 없습니다.', 'error')
        return redirect(url_for('index'))
    
    if request.method == 'POST':
        hire_date = request.form.get('hire_date')
        
        # 입사일로부터 근속연수 자동 재계산
        if hire_date:
            hire_dt = datetime.strptime(hire_date, '%Y-%m-%d')
            years_of_service = relativedelta(datetime.now(), hire_dt).years
        else:
            years_of_service = int(request.form.get('years_of_service') or 0)
        
        # 연차 자동 재계산 (한국 근로기준법)
        annual_leave = calculate_annual_leave_kr(hire_date) if hire_date else int(request.form.get('annual_leave') or 15)
        
        employee['name'] = request.form.get('name')
        employee['ssn'] = request.form.get('ssn')
        employee['department'] = request.form.get('department')
        employee['position'] = request.form.get('position')
        employee['hire_date'] = hire_date
        employee['resignation_date'] = request.form.get('resignation_date') or None
        employee['address'] = request.form.get('address')
        employee['personal_email'] = request.form.get('personal_email')
        employee['work_email'] = request.form.get('work_email')
        employee['personal_phone'] = request.form.get('personal_phone')
        employee['work_phone'] = request.form.get('work_phone')
        employee['has_insurance'] = request.form.get('has_insurance') == 'on'
        employee['salary_bank'] = request.form.get('salary_bank')
        employee['salary_account'] = request.form.get('salary_account')
        employee['years_of_service'] = years_of_service
        employee['annual_leave'] = annual_leave
        employee['used_leave'] = int(request.form.get('used_leave') or 0)
        employee['updated_at'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        
        save_employees(employees)
        flash('직원 정보가 수정되었습니다.', 'success')
        return redirect(url_for('employee_detail', employee_id=employee_id))
    
    return render_template('edit_employee.html', employee=employee)

# 직원 삭제
@app.route('/delete/<int:employee_id>', methods=['POST'])
def delete_employee(employee_id):
    employees = load_employees()
    employees = [emp for emp in employees if emp['id'] != employee_id]
    save_employees(employees)
    
    flash('직원이 삭제되었습니다.', 'success')
    return redirect(url_for('index'))

# 급여 정보 추가/수정
@app.route('/salary/<int:employee_id>', methods=['GET', 'POST'])
def manage_salary(employee_id):
    employees = load_employees()
    employee = next((emp for emp in employees if emp['id'] == employee_id), None)
    
    if employee is None:
        flash('직원을 찾을 수 없습니다.', 'error')
        return redirect(url_for('index'))
    
    salaries = load_salaries()
    
    if request.method == 'POST':
        year = int(request.form.get('year'))
        
        # 해당 직원의 해당 연도 급여 정보 찾기
        salary_entry = next((s for s in salaries if s['employee_id'] == employee_id and s['year'] == year), None)
        
        salary_data = {
            'employee_id': employee_id,
            'year': year,
            'month_01': int(request.form.get('month_01') or 0),
            'month_02': int(request.form.get('month_02') or 0),
            'month_03': int(request.form.get('month_03') or 0),
            'month_04': int(request.form.get('month_04') or 0),
            'month_05': int(request.form.get('month_05') or 0),
            'month_06': int(request.form.get('month_06') or 0),
            'month_07': int(request.form.get('month_07') or 0),
            'month_08': int(request.form.get('month_08') or 0),
            'month_09': int(request.form.get('month_09') or 0),
            'month_10': int(request.form.get('month_10') or 0),
            'month_11': int(request.form.get('month_11') or 0),
            'month_12': int(request.form.get('month_12') or 0),
        }
        
        if salary_entry:
            # 업데이트
            salary_entry.update(salary_data)
        else:
            # 새로 추가
            salaries.append(salary_data)
        
        save_salaries(salaries)
        flash('급여 정보가 저장되었습니다.', 'success')
        return redirect(url_for('employee_detail', employee_id=employee_id, year=year))
    
    # GET 요청 - 급여 입력 폼 표시
    year = request.args.get('year', type=int, default=datetime.now().year)
    salary_info = next((s for s in salaries if s['employee_id'] == employee_id and s['year'] == year), None)
    
    return render_template('manage_salary.html', employee=employee, salary_info=salary_info, year=year)

# 전직원 연차 일괄 재계산
@app.route('/recalculate_annual_leave', methods=['POST'])
def recalculate_annual_leave():
    employees = load_employees()
    updated_count = 0
    
    for employee in employees:
        if employee.get('hire_date'):
            old_leave = employee.get('annual_leave', 0)
            new_leave = calculate_annual_leave_kr(employee['hire_date'])
            
            if old_leave != new_leave:
                employee['annual_leave'] = new_leave
                updated_count += 1
    
    save_employees(employees)
    flash(f'전체 직원 {len(employees)}명의 연차를 재계산했습니다. (변경: {updated_count}명)', 'success')
    return redirect(url_for('index'))

# 엑셀 예제 시트 다운로드
@app.route('/download_excel_template')
def download_excel_template():
    """직원 정보 입력용 엑셀 템플릿 다운로드"""
    wb = Workbook()
    ws = wb.active
    ws.title = "직원정보"
    
    # 헤더 스타일
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    # 헤더 작성
    headers = [
        '이름*', '주민번호*', '부서*', '직급*', '입사일*(YYYY-MM-DD)', 
        '퇴사일(YYYY-MM-DD)', '주소', '개인이메일', '업무이메일',
        '개인전화', '업무전화', '사대보험(O/X)', '급여은행', '급여계좌',
        '사용연차일'
    ]
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 예제 데이터 추가
    example_data = [
        ['홍길동', '900101-1234567', '개발팀', '대리', '2020-01-15', '', 
         '서울시 강남구', 'hong@example.com', 'hong@company.com',
         '010-1234-5678', '02-1234-5678', 'O', '신한은행', '110-123-456789', '3'],
        ['김영희', '850303-2234567', '인사팀', '과장', '2018-03-01', '',
         '서울시 서초구', 'kim@example.com', 'kim@company.com',
         '010-2345-6789', '02-2345-6789', 'O', '국민은행', '123-456-789012', '5']
    ]
    
    for row_num, row_data in enumerate(example_data, 2):
        for col_num, value in enumerate(row_data, 1):
            ws.cell(row=row_num, column=col_num, value=value)
    
    # 열 너비 자동 조정
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # 메모리에 저장
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='직원정보_템플릿.xlsx'
    )

# 엑셀 파일 업로드 및 데이터 import
@app.route('/upload_excel', methods=['GET', 'POST'])
def upload_excel():
    if request.method == 'POST':
        if 'file' not in request.files:
            flash('파일이 선택되지 않았습니다.', 'error')
            return redirect(url_for('upload_excel'))
        
        file = request.files['file']
        
        if file.filename == '':
            flash('파일이 선택되지 않았습니다.', 'error')
            return redirect(url_for('upload_excel'))
        
        if not file.filename.endswith(('.xlsx', '.xls')):
            flash('엑셀 파일(.xlsx, .xls)만 업로드 가능합니다.', 'error')
            return redirect(url_for('upload_excel'))
        
        try:
            # 엑셀 파일 읽기
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            
            employees = load_employees()
            new_id = max([emp.get('id', 0) for emp in employees], default=0) + 1
            
            added_count = 0
            error_rows = []
            
            # 헤더 행 건너뛰고 데이터 읽기
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                # 빈 행 건너뛰기
                if not any(row):
                    continue
                
                try:
                    # 필수 필드 확인
                    name = row[0]
                    ssn = row[1]
                    department = row[2]
                    position = row[3]
                    hire_date = row[4]
                    
                    if not all([name, ssn, department, position, hire_date]):
                        error_rows.append(f"행 {row_num}: 필수 항목 누락")
                        continue
                    
                    # 날짜 형식 변환
                    if isinstance(hire_date, datetime):
                        hire_date = hire_date.strftime('%Y-%m-%d')
                    else:
                        hire_date = str(hire_date)
                    
                    resignation_date = row[5]
                    if resignation_date and isinstance(resignation_date, datetime):
                        resignation_date = resignation_date.strftime('%Y-%m-%d')
                    elif resignation_date:
                        resignation_date = str(resignation_date)
                    else:
                        resignation_date = None
                    
                    # 사대보험 여부 변환
                    insurance_value = str(row[11]).strip().upper() if row[11] else ''
                    has_insurance = insurance_value in ['O', 'Y', 'YES', '1', 'TRUE']
                    
                    # 연차 자동 계산
                    annual_leave = calculate_annual_leave_kr(hire_date)
                    used_leave = int(row[14]) if row[14] and str(row[14]).strip() else 0
                    
                    # 근속연수 계산
                    hire_dt = datetime.strptime(hire_date, '%Y-%m-%d')
                    years_of_service = relativedelta(datetime.now(), hire_dt).years
                    
                    # 새 직원 정보 생성
                    new_employee = {
                        'id': new_id,
                        'name': str(name),
                        'ssn': str(ssn),
                        'department': str(department),
                        'position': str(position),
                        'hire_date': hire_date,
                        'resignation_date': resignation_date,
                        'address': str(row[6]) if row[6] else '',
                        'personal_email': str(row[7]) if row[7] else '',
                        'work_email': str(row[8]) if row[8] else '',
                        'personal_phone': str(row[9]) if row[9] else '',
                        'work_phone': str(row[10]) if row[10] else '',
                        'has_insurance': has_insurance,
                        'salary_bank': str(row[12]) if row[12] else '',
                        'salary_account': str(row[13]) if row[13] else '',
                        'years_of_service': years_of_service,
                        'annual_leave': annual_leave,
                        'used_leave': used_leave,
                        'created_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                    }
                    
                    employees.append(new_employee)
                    new_id += 1
                    added_count += 1
                    
                except Exception as e:
                    error_rows.append(f"행 {row_num}: {str(e)}")
            
            # 저장
            save_employees(employees)
            
            # 결과 메시지
            if added_count > 0:
                flash(f'{added_count}명의 직원이 추가되었습니다.', 'success')
            
            if error_rows:
                error_msg = '<br>'.join(error_rows[:10])  # 최대 10개만 표시
                if len(error_rows) > 10:
                    error_msg += f'<br>...외 {len(error_rows) - 10}개'
                flash(f'오류 발생:<br>{error_msg}', 'error')
            
            return redirect(url_for('index'))
            
        except Exception as e:
            flash(f'파일 처리 중 오류 발생: {str(e)}', 'error')
            return redirect(url_for('upload_excel'))
    
    # GET 요청 - 업로드 폼 표시
    return render_template('upload_excel.html')

if __name__ == '__main__':
    # Debug mode should only be enabled in development
    # Set environment variable FLASK_DEBUG=1 for development
    import os
    debug_mode = os.environ.get('FLASK_DEBUG', '0') == '1'
    app.run(debug=debug_mode, host='0.0.0.0', port=5000)
